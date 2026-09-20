"""TXT/Markdown/DOCX parsing, preserving section headings as boundaries
(spec: "File parsing" -> DOCX/Markdown/Text).
"""

import re
from io import BytesIO

from app.extractors.parser import ParsedDocument, ParsedPage

_MARKDOWN_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$", re.MULTILINE)


def parse_text(content: str) -> ParsedDocument:
    headings = list(_MARKDOWN_HEADING_RE.finditer(content))
    if not headings:
        return ParsedDocument(
            pages=[ParsedPage(page_number=None, text=content)], raw_text=content
        )

    pages: list[ParsedPage] = []
    for index, match in enumerate(headings):
        start = match.start()
        end = headings[index + 1].start() if index + 1 < len(headings) else len(content)
        pages.append(
            ParsedPage(
                page_number=None,
                text=content[start:end].strip(),
                section_title=match.group(2).strip(),
            )
        )
    return ParsedDocument(pages=pages, raw_text=content)


def parse_docx(content_bytes: bytes) -> ParsedDocument:
    from docx import Document

    document = Document(BytesIO(content_bytes))
    pages: list[ParsedPage] = []
    current_section: str | None = None
    current_lines: list[str] = []

    def flush() -> None:
        if current_lines:
            pages.append(
                ParsedPage(
                    page_number=None,
                    text="\n".join(current_lines).strip(),
                    section_title=current_section,
                )
            )

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if paragraph.style.name.lower().startswith("heading"):
            flush()
            current_section = text
            current_lines = []
        else:
            current_lines.append(text)
    flush()

    if not pages:
        return ParsedDocument(pages=[], raw_text="")

    raw_text = "\n\n".join(page.text for page in pages)
    return ParsedDocument(pages=pages, raw_text=raw_text)
